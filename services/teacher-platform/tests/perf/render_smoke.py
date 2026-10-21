"""Render Smoke Test（03-ACCEPTANCE §11，M2-B 生死 Gate 配套；每 Gate 回归）。

两层：
- **COM 真机版（Gate 验收必须）**：Windows PowerPoint COM 自动化打开 PPTX → 导出 PNG 逐页
  → 输出 `acceptance/screenshots/m2b-cruel/`（比解析 XML 更接近真实用户）。
- **XML 兜底版（无 Office 的 CI）**：python-pptx 断言——OMML 原生公式存在、字号 ≥20pt、
  无空占位文本。CI 兜底不替代 Gate 验收的 COM 版。

跑法：
    python tests/perf/render_smoke.py <deck.pptx>              # 自动选择 COM 或 XML
    python tests/perf/render_smoke.py <deck.pptx> --com       # 强制 COM（无 Office 报错）
"""
from __future__ import annotations

import argparse
import shutil
import subprocess
import sys
from pathlib import Path

MIN_FONT_PT = 20.0  # L4 判定②：字号 ≥20pt、无内容截断


# ---------------- COM 真机版 ----------------


def render_via_com(pptx: Path, out_dir: Path) -> list[Path]:
    """PowerPoint COM 打开并导出 PNG。返回逐页 PNG 路径（顺序=页序）。"""
    try:
        import win32com.client  # noqa: F401  pywin32
    except ImportError:
        raise RuntimeError(
            "pywin32 未安装（pip install pywin32）或非 Windows——Gate 验收必须真机 COM 版"
        )
    out_dir.mkdir(parents=True, exist_ok=True)
    ppt_abs = str(pptx.resolve())
    slide_dir = out_dir / "slides"
    if slide_dir.exists():
        shutil.rmtree(slide_dir)
    slide_dir.mkdir(parents=True)
    app = win32com.client.Dispatch("PowerPoint.Application")
    pres = app.Presentations.Open(ppt_abs, ReadOnly=True, WithWindow=False)
    try:
        pres.SaveAs(str(slide_dir.resolve()), 18)  # ppSaveAsPNG = 18
    finally:
        pres.Close()
    app.Quit()
    pngs = sorted(slide_dir.glob("*.png"), key=lambda p: p.stat().st_mtime)
    # 用 PowerPoint 自身导 PDF 供人工比对（可选失败不阻塞）
    try:
        pdf_path = out_dir / (pptx.stem + ".pdf")
        app2 = win32com.client.Dispatch("PowerPoint.Application")
        pres2 = app2.Presentations.Open(ppt_abs, ReadOnly=True, WithWindow=False)
        pres2.SaveAs(str(pdf_path.resolve()), 32)  # ppSaveAsPDF = 32
        pres2.Close()
        app2.Quit()
    except Exception as exc:  # noqa: BLE001
        print(f"[warn] PDF 导出失败（不阻塞）：{exc}", file=sys.stderr)
    return pngs


# ---------------- XML 兜底版（CI） ----------------


def assert_xml_contract(pptx: Path) -> list[str]:
    """python-pptx 断言：OMML 存在、字号 ≥20pt（显式设置时）、无空占位。返回问题列表。"""
    problems: list[str] = []
    try:
        from pptx import Presentation
        from pptx.util import Pt
    except ImportError:
        return ["python-pptx 未安装——XML 兜底不可用（pip install python-pptx）"]

    prs = Presentation(str(pptx))
    omml_slides = 0
    for idx, slide in enumerate(prs.slides, start=1):
        slide_xml = slide._element.xml
        if "m:oMath" in slide_xml:
            omml_slides += 1
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if run.font.size is not None and run.font.size < Pt(MIN_FONT_PT):
                        problems.append(
                            f"页 {idx}: 字号 {run.font.size.pt:.1f}pt < {MIN_FONT_PT}pt（{run.text[:20]!r}）"
                        )
            text = shape.text_frame.text.strip()
            if text in ("", " ", "TODO", "占位"):
                problems.append(f"页 {idx}: 空占位文本框")
    if omml_slides == 0:
        problems.append("全篇未检出 m:oMath——公式可能被转图片（M2-B 直接 FAIL 判据：OMML 可编辑）")
    return problems


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx", type=Path)
    ap.add_argument("--com", action="store_true", help="强制 COM 版（Gate 验收）")
    ap.add_argument("--out", type=Path, default=Path("acceptance/screenshots/m2b-cruel"))
    args = ap.parse_args()

    pptx: Path = args.pptx
    if not pptx.exists():
        print(f"[fail] PPTX 不存在：{pptx}", file=sys.stderr)
        return 2

    use_com = args.com or shutil.which("powerpnt") or _has_office()
    if use_com:
        try:
            pngs = render_via_com(pptx, args.out)
            print(f"[COM] 导出 {len(pngs)} 页 PNG → {args.out}")
            problems = assert_xml_contract(pptx)
            if problems:
                print("[XML 契约问题]")
                for p in problems:
                    print("  -", p)
                return 1
            # 逐页图像检查钩子（M2-B 接像素级 diff/空白页检测）
            print("[COM + XML] Render Smoke PASS（逐页图像人工比对入 acceptance/screenshots/m2b-cruel/）")
            return 0
        except RuntimeError as exc:
            if args.com:
                print(f"[fail] {exc}", file=sys.stderr)
                return 3
            print(f"[warn] {exc}——降级 XML 兜底", file=sys.stderr)

    problems = assert_xml_contract(pptx)
    if problems:
        print("[XML 兜底] 契约问题（CI 档）：")
        for p in problems:
            print("  -", p)
        return 1
    print("[XML 兜底] PASS（注意：Gate 验收必须 Windows 真机 COM 版）")
    return 0


def _has_office() -> bool:
    """探测 PowerPoint COM 是否注册。"""
    try:
        import win32com.client

        app = win32com.client.Dispatch("PowerPoint.Application")
        app.Quit()
        return True
    except Exception:  # noqa: BLE001
        return False


if __name__ == "__main__":
    raise SystemExit(main())
