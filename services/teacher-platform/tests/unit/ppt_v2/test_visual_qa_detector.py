"""T24 逐页视觉 QA 检测器测试（合成用例 + TP-H1 final 真实回放）。

为什么合成用例：PPTX 几何 / PNG 像素两条链路必须能在无 COM、无网络、无 DB 的情况下
独立验证——合成产物是唯一能精确构造「检出」与「不误报」边界的办法。
为什么真实回放：TP-H1 final 是人工对照 contact sheet 签过字的基线，检测器必须在它上面
复现 rubric 结论，否则 8100 的 QA 面会把「已验收的 deck」误判为失败。

证据：T24/evidence/qa-detector-RED.txt → qa-detector-GREEN.txt；
      真实回放结果 → T24/evidence/qa-detector-tp-h1-final.json。
"""

from __future__ import annotations

import json
import os
import re
from pathlib import Path

import pytest
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from app.ppt_v2.qa import (
    DETECTOR_VERSION,
    INFO_KINDS,
    P0_KINDS,
    P1_KINDS,
    RUBRIC,
    detect_deck_qa,
    detect_page_qa,
)
from app.ppt_v2.qa.detector import FONT_CHAIN, SCHEMA, missing_glyphs

CONTROL_ROOT = Path(r"D:\ppt-rescue-control\20260911-math-ai-platform-vnext")
TP_H1_FINAL = CONTROL_ROOT / "TP-H1" / "final"
EVIDENCE_DIR = CONTROL_ROOT / "T24" / "evidence"
_SERVICES_ROOT = Path(__file__).resolve().parents[4]
ENGINE_FONTS_PY = _SERVICES_ROOT / "ppt-presenton" / "presenton_adapter" / "engine" / "fonts.py"

SLIDE_W_IN = 13.333
SLIDE_H_IN = 7.5


# ─────────────────────────────────────────────────────────────────────────────
# 合成产物构造（python-pptx + Pillow，零外部依赖）
# ─────────────────────────────────────────────────────────────────────────────


def _deck() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W_IN)
    prs.slide_height = Inches(SLIDE_H_IN)
    return prs


def _slide(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])  # 6 = Blank，无占位符干扰


def _text(slide, name: str, x, y, w, h, text: str, *, size=18.0, color="1E293B"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.name = name
    run = box.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    return box


def _rect(slide, name: str, x, y, w, h, fill: str):
    box = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))  # 1 = RECTANGLE
    box.name = name
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor.from_string(fill)
    return box


def _render(path: Path, *, mode: str = "normal") -> Path:
    """造对应渲染 PNG：normal=正常版面；low_ink=ink<6%；blank=全白。"""
    im = Image.new("RGB", (1280, 720), (255, 255, 255))
    if mode != "blank":
        draw = ImageDraw.Draw(im)
        if mode == "normal":
            for row in range(14):
                top = 60 + row * 42
                draw.rectangle([120, top, 1160, top + 16], fill=(30, 41, 59))
        else:  # low_ink：均匀细线 → ink<6% 但灰度方差远大于 6.0（不与 blank 混淆）
            for row in range(40):
                top = 10 + row * 18
                draw.line([(120, top), (1160, top)], fill=(30, 41, 59), width=1)
    path.parent.mkdir(parents=True, exist_ok=True)
    im.save(path)
    return path


def _finish(tmp_path: Path, prs: Presentation, *, mode: str = "normal", png: bool = True):
    tmp_path.mkdir(parents=True, exist_ok=True)
    pptx = tmp_path / "deck.pptx"
    prs.save(str(pptx))
    pages = tmp_path / "pages"
    if png:
        for page_no in range(1, len(prs.slides) + 1):
            _render(pages / f"page-{page_no:02d}.png", mode=mode)
    return pptx, pages


def _page(report: dict, page_no: int = 1) -> dict:
    return next(p for p in report["pages"] if p["page_no"] == page_no)


def _kinds(report: dict, page_no: int = 1) -> set[str]:
    return {d["kind"] for d in _page(report, page_no)["defects"]}


def _defect(report: dict, kind: str, page_no: int = 1) -> dict:
    hits = [d for d in _page(report, page_no)["defects"] if d["kind"] == kind]
    assert hits, f"page {page_no} 期望检出 {kind}，实际 {_page(report, page_no)['defects']}"
    assert len(hits) == 1, f"单页内 {kind} 应最多 1 条，实际 {len(hits)}"
    return hits[0]


# ─────────────────────────────────────────────────────────────────────────────
# 1. 合成用例：逐个断言检出与不误报
# ─────────────────────────────────────────────────────────────────────────────


def test_overlap_two_text_boxes_detected(tmp_path: Path) -> None:
    prs = _deck()
    slide = _slide(prs)
    _text(slide, "box-a", 1.20, 2.00, 3.00, 1.00, "重叠检测文本框甲")
    _text(slide, "box-b", 2.50, 2.50, 3.00, 1.00, "重叠检测文本框乙")
    pptx, pages = _finish(tmp_path, prs)

    report = detect_deck_qa(pptx, pages)
    defect = _defect(report, "overlap")
    assert defect["severity"] == "p0"
    assert set(defect["shape_ids"]) == {"box-a", "box-b"}
    assert _page(report)["status"] == "fail"


def test_overlap_single_axis_intersection_not_reported(tmp_path: Path) -> None:
    """只有 x 轴相交、y 轴分离 → 不是重叠（TP-H1 要求两轴都 > 0.04in）。"""
    prs = _deck()
    slide = _slide(prs)
    _text(slide, "box-a", 1.20, 2.00, 3.00, 1.00, "同列上段文本内容")
    _text(slide, "box-b", 2.50, 3.50, 3.00, 1.00, "同列下段文本内容")
    pptx, pages = _finish(tmp_path, prs)
    assert "overlap" not in _kinds(detect_deck_qa(pptx, pages))


def test_overflow_beyond_slide_detected(tmp_path: Path) -> None:
    prs = _deck()
    slide = _slide(prs)
    _text(slide, "wide-box", 10.00, 2.00, 4.00, 1.20, "越界文本框内容检查")
    pptx, pages = _finish(tmp_path, prs)

    report = detect_deck_qa(pptx, pages)
    defect = _defect(report, "overflow")
    assert defect["severity"] == "p0"
    assert defect["shape_ids"] == ["wide-box"]
    assert "off_slide" not in _kinds(report)


def test_off_slide_negative_coordinate_detected(tmp_path: Path) -> None:
    prs = _deck()
    slide = _slide(prs)
    _text(slide, "neg-box", -0.60, 3.00, 2.00, 0.80, "负坐标文本框检查")
    pptx, pages = _finish(tmp_path, prs)

    report = detect_deck_qa(pptx, pages)
    defect = _defect(report, "off_slide")
    assert defect["severity"] == "p0"
    assert defect["shape_ids"] == ["neg-box"]
    assert "overflow" not in _kinds(report)


def test_duplicate_long_text_detected_and_short_text_ignored(tmp_path: Path) -> None:
    prs = _deck()
    long_slide = _slide(prs)
    _text(long_slide, "dup-a", 1.20, 2.00, 4.00, 0.60, "重复文本内容一致性检查")
    _text(long_slide, "dup-b", 1.20, 3.20, 4.00, 0.60, "重复文本内容一致性检查")
    short_slide = _slide(prs)
    _text(short_slide, "short-a", 1.20, 2.00, 4.00, 0.60, "短文本")
    _text(short_slide, "short-b", 1.20, 3.20, 4.00, 0.60, "短文本")
    pptx, pages = _finish(tmp_path, prs)

    report = detect_deck_qa(pptx, pages)
    defect = _defect(report, "duplicate", 1)
    assert defect["severity"] == "p0"
    assert set(defect["shape_ids"]) == {"dup-a", "dup-b"}
    # ≤ 6 字符（去空白后）不构成 duplicate（TP-H1 阈值）
    assert "duplicate" not in _kinds(report, 2)


def test_graphic_empty_placeholder_requires_plan_declaration(tmp_path: Path) -> None:
    prs = _deck()
    slide = _slide(prs)
    _rect(slide, "empty-frame", 1.20, 1.60, 5.00, 3.20, "F1F5F9")
    pptx, pages = _finish(tmp_path, prs)
    plan = {"slides": [{"page_no": 1, "slide_id": "slide-01", "asset_bindings": [{"slot": "figure", "asset_id": "a-1"}]}]}

    declared = detect_deck_qa(pptx, pages, plan=plan)
    defect = _defect(declared, "graphic")
    assert defect["severity"] == "p0"
    assert defect["shape_ids"] == ["empty-frame"]

    # plan 未声明资产 → 不得报 graphic（大填色矩形可能是正常留白/底色）
    assert "graphic" not in _kinds(detect_deck_qa(pptx, pages))
    assert "graphic" not in _kinds(detect_deck_qa(pptx, pages, plan={"slides": []}))


def test_contrast_low_ratio_flagged_and_black_on_white_clean(tmp_path: Path) -> None:
    prs = _deck()
    low = _slide(prs)
    _rect(low, "soft-card", 1.20, 1.80, 5.00, 1.60, "E0E0E0")
    _text(low, "gray-on-gray", 1.40, 2.00, 4.40, 0.80, "浅灰字配浅灰底对比度检查", color="888888")
    clean = _slide(prs)
    _text(clean, "black-on-white", 1.20, 2.00, 6.00, 0.80, "黑字白底不构成对比度缺陷", color="000000")
    pptx, pages = _finish(tmp_path, prs)

    report = detect_deck_qa(pptx, pages)
    defect = _defect(report, "contrast", 1)
    assert defect["severity"] == "p0"
    assert defect["shape_ids"] == ["gray-on-gray"]
    assert "contrast" not in _kinds(report, 2)


def test_font_tiny_body_flagged(tmp_path: Path) -> None:
    prs = _deck()
    slide = _slide(prs)
    _text(slide, "tiny-body", 1.20, 2.00, 6.00, 0.80, "八字号正文字号检查", size=8.0)
    pptx, pages = _finish(tmp_path, prs)

    defect = _defect(detect_deck_qa(pptx, pages), "font")
    assert defect["severity"] == "p1"
    assert defect["shape_ids"] == ["tiny-body"]


def test_font_missing_glyph_flagged(tmp_path: Path) -> None:
    if not _font_chain_available():
        pytest.skip("本机字体链不存在（无 msyh.ttc 等），字形覆盖检查无法判定")
    unsupported = "\U00010350"  # Old Permic：链上 7 个字体均无 cmap 覆盖
    assert unsupported in missing_glyphs(f"缺字形检测{unsupported}")

    prs = _deck()
    slide = _slide(prs)
    _text(slide, "glyph-miss", 1.20, 2.00, 6.00, 0.80, f"缺字形检测{unsupported}")
    pptx, pages = _finish(tmp_path, prs)

    defect = _defect(detect_deck_qa(pptx, pages), "font")
    assert defect["severity"] == "p1"


def _font_chain_available() -> bool:
    root = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts"
    return any((root / filename).exists() for _, filename in FONT_CHAIN)


def test_formula_covered_math_chars_no_defect(tmp_path: Path) -> None:
    if not _font_chain_available():
        pytest.skip("本机字体链不存在，字形覆盖检查无法判定")
    prs = _deck()
    slide = _slide(prs)
    _text(slide, "math-line", 1.20, 2.00, 8.00, 0.80, "含数学字符：a ∈ A 且 x² ≥ 0，A ∩ B = ∅。")
    pptx, pages = _finish(tmp_path, prs)

    kinds = _kinds(detect_deck_qa(pptx, pages))
    assert "formula" not in kinds
    assert "font" not in kinds


def test_formula_truncation_losing_math_chars_flagged(tmp_path: Path) -> None:
    """截断把表达式尾巴吃掉 → 花括号不闭合 + 结尾悬挂算符，数学字符已丢失。"""
    prs = _deck()
    slide = _slide(prs)
    _text(slide, "cut-line", 1.20, 2.00, 8.00, 0.80, "截断样例：A ∩ B = {x | x ∈")
    pptx, pages = _finish(tmp_path, prs)

    defect = _defect(detect_deck_qa(pptx, pages), "formula")
    assert defect["severity"] == "p1"
    assert "截断" in defect["message"]


def test_blank_png_flagged(tmp_path: Path) -> None:
    prs = _deck()
    slide = _slide(prs)
    _text(slide, "title", 1.20, 2.00, 6.00, 0.80, "空白页检测用标题")
    pptx, pages = _finish(tmp_path, prs, mode="blank")

    report = detect_deck_qa(pptx, pages)
    defect = _defect(report, "blank")
    assert defect["severity"] == "p1"
    assert _page(report)["status"] == "fail"


def test_density_is_info_and_never_fails_page(tmp_path: Path) -> None:
    prs = _deck()
    slide = _slide(prs)
    _text(slide, "title", 1.20, 2.00, 6.00, 0.80, "低墨量页检测用标题")
    pptx, pages = _finish(tmp_path, prs, mode="low_ink")

    report = detect_deck_qa(pptx, pages)
    defect = _defect(report, "density")
    assert defect["severity"] == "info"
    assert _page(report)["status"] == "pass"
    assert report["summary"]["p0_count"] == 0
    assert report["summary"]["p1_count"] == 0
    assert _page(report)["metrics"]["ink_pct"] < 6.0


def test_missing_png_is_unknown_but_pptx_defect_still_fails(tmp_path: Path) -> None:
    clean = _deck()
    _text(_slide(clean), "title", 1.20, 2.00, 6.00, 0.80, "无渲染证据的干净页")
    clean_pptx, clean_pages = _finish(tmp_path / "clean", clean, png=False)
    assert _page(detect_deck_qa(clean_pptx, clean_pages))["status"] == "unknown"

    bad = _deck()
    _text(_slide(bad), "wide-box", 10.00, 2.00, 4.00, 1.20, "无渲染证据的越界页")
    bad_pptx, bad_pages = _finish(tmp_path / "bad", bad, png=False)
    assert _page(detect_deck_qa(bad_pptx, bad_pages))["status"] == "fail"


def test_detect_page_qa_element_shape_matches_deck_page(tmp_path: Path) -> None:
    prs = _deck()
    slide = _slide(prs)
    _text(slide, "box-a", 1.20, 2.00, 3.00, 1.00, "重叠检测文本框甲")
    _text(slide, "box-b", 2.50, 2.50, 3.00, 1.00, "重叠检测文本框乙")
    pptx, pages = _finish(tmp_path, prs)
    png = pages / "page-01.png"

    single = detect_page_qa(pptx, 1, png)
    assert {"schema", "detector_version"} <= set(single)
    element = {k: v for k, v in single.items() if k not in ("schema", "detector_version")}
    assert element == _page(detect_deck_qa(pptx, pages))
    assert single["schema"] == SCHEMA and single["detector_version"] == DETECTOR_VERSION

    with pytest.raises(ValueError):
        detect_page_qa(pptx, 99, png)


# ─────────────────────────────────────────────────────────────────────────────
# 2. 冻结接口与契约形状
# ─────────────────────────────────────────────────────────────────────────────


def test_frozen_interface_constants() -> None:
    assert DETECTOR_VERSION == "t24-visual-qa-1.0.0"
    assert RUBRIC == "tp-h1-visual-rubric"
    assert SCHEMA == "tp-h1-visual-qa/v1"
    assert P0_KINDS == ("overflow", "overlap", "off_slide", "duplicate", "graphic", "contrast")
    assert P1_KINDS == ("font", "formula")
    assert INFO_KINDS == ("density",)


def test_report_envelope_and_metrics_shape(tmp_path: Path) -> None:
    prs = _deck()
    slide = _slide(prs)
    _text(slide, "box-a", 1.20, 2.00, 3.00, 1.00, "报告信封字段检查甲")
    pptx, pages = _finish(tmp_path, prs)

    report = detect_deck_qa(pptx, pages)
    assert set(report) == {"schema", "detector_version", "pages", "summary"}
    assert set(report["summary"]) == {"p0_count", "p1_count", "failed_pages", "by_kind"}
    assert set(report["pages"][0]) == {"page_no", "status", "defects", "metrics"}
    assert set(report["pages"][0]["metrics"]) == {"ink_pct", "cv", "shape_count", "has_picture"}
    for page in report["pages"]:
        for defect in page["defects"]:
            assert set(defect) == {"kind", "severity", "message", "shape_ids"}
            assert defect["severity"] in ("p0", "p1", "info")
            assert isinstance(defect["shape_ids"], list)
            assert all(isinstance(s, str) for s in defect["shape_ids"])


def test_font_chain_mirrors_presenton_engine() -> None:
    """镜像链的漂移守卫：8100 侧检测器只镜像链序，不得 import 8201 生成侧代码。"""
    if not ENGINE_FONTS_PY.exists():  # pragma: no cover - 引擎源码不在本机
        pytest.skip(f"引擎字体链源文件不存在：{ENGINE_FONTS_PY}")
    text = ENGINE_FONTS_PY.read_text(encoding="utf-8")
    block = re.search(r"FONT_CHAIN[^=]*=\s*\((.*?)\n\)", text, re.S)
    assert block, "未能在 engine/fonts.py 中定位 FONT_CHAIN"
    engine_chain = tuple(re.findall(r'\("([^"]+)",\s*"([^"]+)"\)', block.group(1)))
    assert engine_chain == FONT_CHAIN


# ─────────────────────────────────────────────────────────────────────────────
# 3. 真实回放：TP-H1 final（12 页真实 PPTX + 真实 COM 渲染 PNG）
# ─────────────────────────────────────────────────────────────────────────────

TP_H1_DECK = TP_H1_FINAL / "deck.pptx"
TP_H1_PAGES = TP_H1_FINAL / "pages"


@pytest.mark.skipif(not TP_H1_DECK.exists(), reason=f"TP-H1 final 不存在：{TP_H1_DECK}")
def test_tp_h1_final_replay_matches_rubric() -> None:
    plan = json.loads((TP_H1_FINAL / "plan.json").read_text(encoding="utf-8"))
    report = detect_deck_qa(TP_H1_DECK, TP_H1_PAGES, plan=plan)

    EVIDENCE_DIR.mkdir(parents=True, exist_ok=True)
    (EVIDENCE_DIR / "qa-detector-tp-h1-final.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    assert len(report["pages"]) == 12
    assert [p["page_no"] for p in report["pages"]] == list(range(1, 13))
    assert report["schema"] == SCHEMA and report["detector_version"] == DETECTOR_VERSION

    # TP-H1 rubric 自评 final = PASS（P0=0 / P1=0，density_low 当时记 P1 但不卡门）。
    # 本检测器对 TP-H1 原有维度必须逐维复现 0：
    #   P0-1 overlap / P0-2 overflow / P0-3 off_slide / P0-4 duplicate / P0-5 graphic
    #   P1-1 missing_glyph→font / P1-2 tiny_font→font / P1-4 page_blank→blank
    tp_h1_dims = {"overflow", "overlap", "off_slide", "duplicate", "graphic", "font", "formula", "blank"}
    assert not (set(report["summary"]["by_kind"]) & tp_h1_dims), report["summary"]["by_kind"]

    # 但本轮新增的 contrast 维（TP-H1 rubric 无此维）在 page-01 命中 1 处真实缺陷：
    #   meta 行「教材直连 · SourceLocator 溯源见备注」= #94A3B8 on #FFFFFF → 2.56:1 < 3.0。
    # 已用渲染 PNG 逐像素复核（meta 区域主色 = (148,163,184) on (255,255,255)）。
    # 阈值来自冻结 rubric，**不因回放结果调整** → 真实 P0 = 1（不是 0）。
    assert report["summary"]["by_kind"].get("contrast") == 1
    assert report["summary"]["p0_count"] == 1
    assert report["summary"]["p1_count"] == 0
    assert report["summary"]["failed_pages"] == [1]

    contrast = _defect(report, "contrast", 1)
    assert contrast["severity"] == "p0"
    assert contrast["shape_ids"] == ["TextBox 11"]

    # density 共 10 页（TP-H1 记为 density_low），按冻结契约降为 info：不卡门、不计入 p1
    assert report["summary"]["by_kind"].get("density") == 10
    assert [p["page_no"] for p in report["pages"] if p["status"] == "pass"] == [2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12]
