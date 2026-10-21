"""T24（8201）：POST /api/v1/deck/export/pdf —— 真实 PowerPoint COM 导出 PDF。

契约：docs/compose/spec/t24-artifact-edit-export.md S2.6 —— 8100 从 8200 取到 PPTX 字节后，
送 8201 → PowerPoint COM `SaveAs(path, 32)` 得到**真实** PDF；失败必须诚实失败（不产出假 PDF）。

页数验证路径（本机实测取舍，2026-09-13）：
- LibreOffice（soffice）本机不存在；
- 8201 venv 无 pypdfium2 / pypdf / PyMuPDF（实测 ModuleNotFoundError）；
- PowerPoint COM 重新打开导出的 PDF 在本机**会挂死**（实测 150s 未返回：PDF reflow 转换阻塞）。
故页数由测试内置的 PDF 结构解析（逐个 /Type /Page 对象计数 + 页树 /Count 交叉核对）
在真实 PDF 字节上验证，并与端点返回的 page_count（来自 PowerPoint `Slides.Count`）对齐。
"""
from __future__ import annotations

import base64
import hashlib
import io
import os
import re
import zlib
from pathlib import Path

import httpx
import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from presenton_adapter.app import create_app
from presenton_adapter.config import get_settings
from presenton_adapter.engine import pdf_export

APP = create_app()
BASE = "http://presenton.test"
TOKEN_HEADER = "X-Presenton-Service-Token"
URL = "/api/v1/deck/export/pdf"

_PAGE_OBJ_RE = re.compile(rb"/Type\s*/Page(?![A-Za-z])")
_COUNT_RE = re.compile(rb"/Count\s+(\d+)")


def _com_available() -> bool:
    """PowerPoint COM 可用性探测（只做探测，不启动 COM，避免挂死影响其余用例）。"""
    try:
        import pythoncom  # noqa: F401
        import win32com.client  # noqa: F401
    except ImportError:
        return False
    try:
        import winreg

        with winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, "PowerPoint.Application"):
            return True
    except OSError:
        return False


COM_AVAILABLE = _com_available()


def _headers() -> dict[str, str]:
    return {TOKEN_HEADER: get_settings().service_token}


def _two_page_pptx() -> bytes:
    """python-pptx 造 2 页真实 PPTX（含中文文本），并自证可被 python-pptx 打开。"""
    prs = Presentation()
    prs.slide_width = int(13.333 * 914400)
    prs.slide_height = int(7.5 * 914400)
    for title, body in [
        ("集合与元素", "集合指一定范围内确定的、不同的对象的全体。\n- 元素与集合的关系用 ∈ 表示"),
        ("一元二次方程", "x² − 3x + 2 = 0 的实数根组成的集合"),
    ]:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        tb = slide.shapes.add_textbox(Inches(0.7), Inches(0.6), Inches(11.0), Inches(1.2))
        tb.text_frame.text = title
        tb.text_frame.paragraphs[0].runs[0].font.size = Pt(32)
        bb = slide.shapes.add_textbox(Inches(0.7), Inches(2.2), Inches(11.0), Inches(4.0))
        bb.text_frame.text = body
        bb.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
    buf = io.BytesIO()
    prs.save(buf)
    data = buf.getvalue()
    assert len(Presentation(io.BytesIO(data)).slides._sldIdLst) == 2
    return data


def _pdf_pages(data: bytes) -> tuple[int, int]:
    """真实 PDF 字节的结构页数：返回 (对象计数, 页树 /Count)。"""
    n = len(_PAGE_OBJ_RE.findall(data))
    for m in re.finditer(rb"stream\r?\n", data):
        start = m.end()
        end = data.find(b"endstream", start)
        if end < 0:
            continue
        try:
            dec = zlib.decompress(data[start:end])
        except Exception:  # noqa: BLE001  非 flate 流/损坏流：跳过（不掩盖计数）
            continue
        n += len(_PAGE_OBJ_RE.findall(dec))
    counts = [int(c) for c in _COUNT_RE.findall(data)]
    return n, (max(counts) if counts else 0)


async def _post(payload: dict, headers: dict[str, str] | None) -> httpx.Response:
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=APP), base_url=BASE) as client:
        return await client.post(URL, json=payload, headers=headers or {})


# ── 正例：真实 COM 导出 ────────────────────────────────────────────


@pytest.mark.skipif(not COM_AVAILABLE, reason="本机无 PowerPoint COM（pywin32/ProgID 不可用）")
async def test_export_pdf_real_two_page_deck() -> None:
    pptx = _two_page_pptx()
    r = await _post({"pptx_b64": base64.b64encode(pptx).decode(), "file_stem": "t24-deck"}, _headers())
    assert r.status_code == 200, r.text
    body = r.json()
    pdf = base64.b64decode(body["pdf_b64"])
    assert pdf[:5] == b"%PDF-", f"非 PDF 字节: {pdf[:8]!r}"
    assert body["bytes"] == len(pdf) > 0
    assert hashlib.sha256(pdf).hexdigest() == body["sha256"]
    assert body["page_count"] == 2
    assert body["engine"]["name"] == "powerpoint-com"
    assert str(body["engine"]["version"]).startswith("16"), body["engine"]
    # 独立结构核对：对象级页数 + 页树 /Count 都是 2（与端点 page_count 一致）
    obj_pages, tree_count = _pdf_pages(pdf)
    assert obj_pages == 2, f"PDF 对象级页数={obj_pages}"
    assert tree_count == 2, f"PDF 页树 /Count={tree_count}"
    ev = os.environ.get("T24_PDF_EVIDENCE_DIR")
    if ev:
        out = Path(ev)
        out.mkdir(parents=True, exist_ok=True)
        (out / "engine-8201-real-export.pdf").write_bytes(pdf)
        (out / "engine-8201-real-export.json").write_text(
            f"page_count={body['page_count']} bytes={body['bytes']} sha256={body['sha256']} "
            f"engine={body['engine']}\nobj_pages={obj_pages} tree_count={tree_count}\n",
            encoding="utf-8",
        )


# ── 鉴权 / 输入负例 ──────────────────────────────────────────────


async def test_export_pdf_missing_token_401() -> None:
    r = await _post({"pptx_b64": base64.b64encode(_two_page_pptx()).decode()}, None)
    assert r.status_code == 401, r.text
    assert r.json()["error_key"] == "unauthorized"


async def test_export_pdf_wrong_token_401() -> None:
    r = await _post({"pptx_b64": base64.b64encode(_two_page_pptx()).decode()}, {TOKEN_HEADER: "wrong"})
    assert r.status_code == 401, r.text


async def test_export_pdf_malformed_base64_4xx() -> None:
    r = await _post({"pptx_b64": "!!!not-base64!!!", "file_stem": "deck"}, _headers())
    assert 400 <= r.status_code < 500, r.text
    assert r.status_code == 422
    assert "pdf_b64" not in r.text


async def test_export_pdf_non_ooxml_bytes_4xx() -> None:
    r = await _post({"pptx_b64": base64.b64encode(b"this is not a pptx").decode()}, _headers())
    assert r.status_code == 422, r.text
    assert "pdf_b64" not in r.text


# ── 诚实失败：COM 不可用绝不返回假 PDF ───────────────────────────


async def test_export_pdf_com_unavailable_fails_honestly(monkeypatch) -> None:  # noqa: ANN001
    def _boom(*_a, **_kw):  # noqa: ANN002, ANN003
        raise pdf_export.PdfExportUnavailable("simulated: PowerPoint COM 不可用")

    monkeypatch.setattr(pdf_export, "_export_pdf_sync", _boom)
    r = await _post({"pptx_b64": base64.b64encode(_two_page_pptx()).decode()}, _headers())
    assert r.status_code == 503, r.text
    assert "pdf_b64" not in r.text
    assert "powerpoint-com" in r.text
