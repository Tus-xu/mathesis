"""最终装配：页 fragment（单页 PPTX）→ 完整可编辑 PPTX + PowerPoint 打开检查 + contact sheet。

拷贝策略：受控 builder 产物 → deepcopy spTree 形状元素 + 背景 + 备注；
图片形状按 blob 重新入库并改写 r:embed（包内关系重建，避免悬空 rId）。
幂等：deck_key（builder_version+spec_hash+page_refs）命中即直接返回已有产物。

R3d 性能回炉（去重复渲染）：逐页 QA 已经用 PowerPoint 渲染过每一页 PNG；
装配不再把整册**再渲染一遍**做 contact sheet —— 直接用逐页真实页图拼图，
最终册只做一次 PowerPoint **打开检查**（openable，真实可打开证明）。页图缺失时
才回退整册渲染（如实降级，不伪造）。
"""
from __future__ import annotations

import copy
import hashlib
import io
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.oxml.ns import qn

from presenton_adapter.engine import renderer
from presenton_adapter.engine.builder import BUILDER_VERSION, SLIDE_H_IN, SLIDE_W_IN

A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


@dataclass
class AssembleResult:
    deck_key: str
    deck_ref: str  # 内容寻址 ref（= sha256(deck bytes)），artifact_bytes 可直接取用
    deck_sha256: str
    size_bytes: int
    page_count: int
    contact_ref: str
    openable: dict
    elapsed_ms: int = 0
    errors: list[str] = field(default_factory=list)


def _copy_slide(src_slide, dst_slide) -> None:
    # 1) 背景
    src_cSld = src_slide._element.cSld
    bg = src_cSld.find(qn("p:bg"))
    if bg is not None:
        dst_slide._element.cSld.insert(0, copy.deepcopy(bg))
    # 2) 形状树
    dst_spTree = dst_slide.shapes._spTree
    pic_fixups: list[tuple[Any, Any]] = []  # (copied pic element, blob)
    for shape in src_slide.shapes:
        el = copy.deepcopy(shape._element)
        if el.tag == qn("p:pic"):
            blip = el.find(f".//{A_NS}blip")
            rid = blip.get(f"{R_NS}embed") if blip is not None else None
            blob = None
            if rid:
                part = src_slide.part.related_part(rid)
                blob = part.blob
            pic_fixups.append((el, blob))
        dst_spTree.append(el)
    # 3) 图片关系重建
    for el, blob in pic_fixups:
        blip = el.find(f".//{A_NS}blip")
        if blip is None:
            continue
        # get_or_add_image_part 返回 (image_part, rId) 元组
        _image_part, new_rid = dst_slide.part.get_or_add_image_part(io.BytesIO(blob or b""))
        blip.set(f"{R_NS}embed", new_rid)
    # 4) 备注
    notes = ""
    if src_slide.has_notes_slide:
        notes = src_slide.notes_slide.notes_text_frame.text or ""
    if notes:
        dst_slide.notes_slide.notes_text_frame.text = notes


def _contact_sheet(png_blobs: list[bytes], out_path: Path, *, cols: int = 4, thumb_w: int = 460) -> Path:
    thumbs = []
    for blob in png_blobs:
        im = Image.open(io.BytesIO(blob)).convert("RGB")
        ratio = thumb_w / im.width
        thumbs.append(im.resize((thumb_w, int(im.height * ratio))))
    if not thumbs:
        raise RuntimeError("contact sheet 无输入页")
    th_h = thumbs[0].height
    label_h = 34
    rows = (len(thumbs) + cols - 1) // cols
    pad = 14
    W = cols * thumb_w + (cols + 1) * pad
    H = rows * (th_h + label_h) + (rows + 1) * pad
    sheet = Image.new("RGB", (W, H), "#F8FAFC")
    draw = ImageDraw.Draw(sheet)
    try:
        font = ImageFont.truetype(r"C:\Windows\Fonts\msyh.ttc", 18)
    except Exception:  # noqa: BLE001
        font = ImageFont.load_default()
    for i, im in enumerate(thumbs):
        r, c = divmod(i, cols)
        x = pad + c * (thumb_w + pad)
        y = pad + r * (th_h + label_h + pad)
        sheet.paste(im, (x, y))
        draw.rectangle([x, y, x + thumb_w, y + th_h], outline="#CBD5E1", width=1)
        draw.text((x + 2, y + th_h + 6), f"第 {i + 1:02d} 页", fill="#1E293B", font=font)
    sheet.save(out_path, "PNG")
    return out_path


def _assemble_sync(store, deck_key: str, page_refs: list[str], out_path: Path) -> None:
    """python-pptx 装配（纯字节拼装，不驱动 PowerPoint）。"""
    prs = Presentation()
    prs.slide_width = int(SLIDE_W_IN * 914400)
    prs.slide_height = int(SLIDE_H_IN * 914400)
    for ref in page_refs:
        frag_bytes = store.get(ref)
        if frag_bytes is None:
            raise RuntimeError(f"page fragment 缺失: {ref[:12]}…")
        frag_path = store.get_path(frag_bytes, "pptx")
        frag_prs = Presentation(str(frag_path))
        blank = prs.slide_layouts[6]
        dst = prs.slides.add_slide(blank)
        _copy_slide(frag_prs.slides[0], dst)
    prs.core_properties.title = "Teaching Deck (R21)"
    prs.core_properties.author = "ppt-studio-v2"
    prs.core_properties.comments = f"builder={BUILDER_VERSION}; deck_key={deck_key}"
    prs.save(str(out_path))


async def assemble_deck(
    store,
    *,
    deck_id: str,
    deck_title: str,
    deck_spec_sha256: str,
    page_refs: list[str],
    page_nos: list[int],
    budget_timeout_s: float = 420.0,
    page_png_bytes: list[bytes] | None = None,
) -> AssembleResult:
    import asyncio

    t0 = time.monotonic()
    deck_key = ArtifactStore_deck_key(deck_spec_sha256, page_refs)
    out_path = store.deck_path(deck_key)
    errors: list[str] = []
    if out_path.exists():
        # 幂等：同请求（同 spec+同页集合）→ 复用，不重建
        data = out_path.read_bytes()
        deck_ref = store.put(data, "pptx")
        n, _, openable = await asyncio.to_thread(_open_only, out_path)
        contact_ref = _contact_ref_for(store, deck_key, data)
        return AssembleResult(
            deck_key=deck_key,
            deck_ref=deck_ref,
            deck_sha256=ArtifactStore_sha(data),
            size_bytes=len(data),
            page_count=n,
            contact_ref=contact_ref,
            openable=openable,
            elapsed_ms=int((time.monotonic() - t0) * 1000),
            errors=errors,
        )
    import tempfile

    with tempfile.TemporaryDirectory(prefix="ps8201-asm-") as td:
        tmp_dir = Path(td)
        await asyncio.wait_for(
            asyncio.to_thread(_assemble_sync, store, deck_key, page_refs, out_path),
            timeout=budget_timeout_s,
        )
        # 真实 PowerPoint 打开检查（最终册唯一 COM 步骤：可打开 + 页数核对）
        openable = await asyncio.to_thread(renderer._openable_sync, out_path)
        n = int(openable.get("slide_count") or 0)
        if n != len(page_refs):
            errors.append(f"page-count-mismatch:{n}!={len(page_refs)}")
        # contact sheet：优先复用逐页 QA 真实页图（去掉整册重复渲染）；缺位才回退整册渲染
        blobs = [b for b in (page_png_bytes or []) if b]
        if len(blobs) == len(page_refs):
            openable = {**openable, "rendered_pngs": 0, "contact_from": "page_pngs"}
        else:
            _n2, pngs = await asyncio.to_thread(renderer._export_sync, out_path, tmp_dir)
            blobs = [Path(p).read_bytes() for p in pngs]
            openable = {**openable, "rendered_pngs": len(pngs), "contact_from": "deck_render"}
        sheet_path = tmp_dir / "contact-sheet.png"
        _contact_sheet(blobs, sheet_path)
        sheet_bytes = sheet_path.read_bytes()
        contact_ref = store.put(sheet_bytes, "png")
        # contact sheet 落 deck 目录旁证（可复现）
        sheet_copy = store.root / "decks" / f"{deck_key}-contact-sheet.png"
        if not sheet_copy.exists():
            sheet_copy.write_bytes(sheet_bytes)
    data = out_path.read_bytes()
    if len(data) <= 0:
        errors.append("artifact-zero-size")
    deck_ref = store.put(data, "pptx")
    return AssembleResult(
        deck_key=deck_key,
        deck_ref=deck_ref,
        deck_sha256=ArtifactStore_sha(data),
        size_bytes=len(data),
        page_count=n,
        contact_ref=contact_ref,
        openable=openable,
        elapsed_ms=int((time.monotonic() - t0) * 1000),
        errors=errors,
    )


def _open_only(pptx_path: Path) -> tuple[int, list[str], dict]:
    res = renderer._openable_sync(pptx_path)
    openable = {**res, "rendered_pngs": 0}
    return int(res.get("slide_count", 0)), [], openable


def ArtifactStore_sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def ArtifactStore_deck_key(deck_spec_sha256: str, page_refs: list[str]) -> str:
    import json

    canon = json.dumps({"v": BUILDER_VERSION, "spec": deck_spec_sha256, "pages": page_refs},
                       sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(canon.encode("utf-8")).hexdigest()


def _contact_ref_for(store, deck_key: str, deck_bytes: bytes) -> str:
    sheet_copy = store.root / "decks" / f"{deck_key}-contact-sheet.png"
    if sheet_copy.exists():
        return store.put(sheet_copy.read_bytes(), "png")
    return ""
