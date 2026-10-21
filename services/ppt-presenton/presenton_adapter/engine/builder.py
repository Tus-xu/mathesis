"""python-pptx 逐页 OOXML 构建（可编辑：真实文本框/形状/图片/备注）。

设计要点：
1. 输入只有语义信息（layout_hint / content / assets）——引擎内部把语义布局
   翻译为形状几何；TeachingDeckSpec 永不出现引擎像素坐标（DECISIONS/AGENTS 铁律）。
2. 每页 = 单页 fragment PPTX（可独立打开、可编辑），装配阶段按页拷贝进最终文档。
3. 内容寻址：content_sha256 = sha256(canonical(builder_version + deck 标题 + 页号 + slide + 资产哈希))
   → 相同输入必得相同键（同请求无重复 artifact 的基础）。
4. 构造期即用 clamp_lines_to_fit 做确定性防溢出；QA 仍独立复核（双保险，构造不掩盖 QA）。
5. R22 皮肤（golden-B-r211 口径）：plan 路径文字全部内嵌 AUTO_SHAPE（页上零裸 TEXT_BOX），
   正文块叠浅底圆角卡片；卡片 qa_exempt，不进 QA 度量盒。
"""
from __future__ import annotations

import hashlib
import io
import json
from dataclasses import dataclass, field
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from presenton_adapter.engine import metrics as M
from presenton_adapter.engine.layout_resolver import Element
from presenton_adapter.engine.metrics import ShapeBox

BUILDER_VERSION = "r21.1-engine-tp-h1-1.4.0"  # 1.4.0：R22 皮肤移植（文字内嵌 AUTO_SHAPE + 正文卡片化，对齐 golden-B-r211）；1.3.0：参数化 math_object → PNG 内嵌（figure 槽位 + 可见失败）；1.2.0：LaTeX→Unicode 过渡排版（R16 前临时方案）

SLIDE_W_IN = 13.333
SLIDE_H_IN = 7.5

INK = RGBColor(0x1E, 0x29, 0x3B)
MUTED = RGBColor(0x64, 0x74, 0x8B)
SUBTLE = RGBColor(0x94, 0xA3, 0xB8)
BRAND = RGBColor(0x25, 0x63, 0xEB)
BRAND_BG = RGBColor(0xEF, 0xF6, 0xFF)
SURFACE = RGBColor(0xFF, 0xFF, 0xFF)
SOFT_BG = RGBColor(0xF1, 0xF5, 0xF9)
LINE = RGBColor(0xE2, 0xE8, 0xF0)

FONT_PRIMARY = "Microsoft YaHei"
# 中文字体 fallback 链（与 fonts.FONT_CHAIN 对齐）
FONT_FALLBACK_CS = "Cambria Math"
FONT_FALLBACK_EA = "SimSun"

# 内部类型 → 中文标签（basic fallback 唯一合法 chip 文本）
TYPE_ZH_LABEL: dict[str, str] = {
    "title": "封面",
    "introduction": "导入",
    "concept": "概念",
    "definition": "定义",
    "derivation": "推导",
    "worked-example": "例题",
    "comparison": "对比",
    "image-focus": "观察",
    "practice": "练习",
    "inquiry": "探究",
    "summary": "小结",
    "overview": "总览",
}

MATH_CHARS = set("∈∉∋⊆⊇⊂⊃∪∩∅≤≥≠≈±×÷²³∑∫√∞∥⊥∠°′″∀∃→←↔⇒⇔·−")


@dataclass
class PageRequest:
    deck_id: str
    deck_title: str
    page_no: int
    total_pages: int
    slide: dict[str, Any]
    assets_bytes: dict[str, bytes] = field(default_factory=dict)
    visual_plan: dict[str, Any] | None = None      # R21.1：per-slide ResolvedVisualPlan（无坐标）
    deck_style_tokens: dict[str, Any] | None = None  # R21.1：deck 级 style/typography tokens


@dataclass
class BuiltPage:
    fragment_bytes: bytes
    boxes: list[ShapeBox]
    content_sha256: str
    notes_text: str
    math_chars: list[str]
    has_figure: bool


def _is_math_line(line: str) -> bool:
    return any(ch in MATH_CHARS for ch in line)


def _norm_bullets(content: str) -> list[str]:
    lines = []
    for raw in content.split("\n"):
        t = raw.strip()
        if not t:
            continue
        if t.startswith("- ") or t.startswith("• "):
            t = "• " + t[2:]
        elif t.startswith(("#", "##")):
            t = t.lstrip("# ").strip()
        lines.append(t)
    return lines


def _set_bg(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _add_box(slide, x: float, y: float, w: float, h: float, *, fill: RGBColor | None = None,
             line: RGBColor | None = None, shape=MSO_SHAPE.RECTANGLE, radius: float | None = None):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(1)
    sp.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sp.adjustments[0] = radius
        except Exception:  # noqa: BLE001
            pass
    return sp


def _write_runs(tf, text: str, *, size: float, color: RGBColor, bold: bool,
                align, single_line: bool, line_spacing: float) -> None:
    """段落/字体样式写入（TEXT_BOX 与 AUTO_SHAPE 两种载体共用，QA 度量同口径）。"""
    tf.word_wrap = not single_line
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    lines = text.split("\n") if not single_line else [text]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        r = p.add_run()
        r.text = ln
        f = r.font
        f.name = FONT_PRIMARY
        f.size = Pt(size)
        f.bold = bold
        f.color.rgb = color
        # 东亚字体名（ea）同步，保证 PowerPoint 中文按雅黑渲染
        rPr = r._r.get_or_add_rPr()
        ea = rPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
        if ea is None:
            from lxml import etree

            ea = etree.SubElement(rPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
        ea.set("typeface", FONT_PRIMARY)


def _add_text(slide, x: float, y: float, w: float, h: float, text: str, *, size: float,
              color: RGBColor = INK, bold: bool = False, align=PP_ALIGN.LEFT,
              single_line: bool = False, anchor=MSO_ANCHOR.TOP, line_spacing: float = 1.18):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.text_frame.vertical_anchor = anchor
    _write_runs(tb.text_frame, text, size=size, color=color, bold=bold, align=align,
                single_line=single_line, line_spacing=line_spacing)
    return tb


def _add_shape_text(slide, x: float, y: float, w: float, h: float, text: str, *, size: float,
                    color: RGBColor = INK, bold: bool = False, align=PP_ALIGN.LEFT,
                    single_line: bool = False, anchor=MSO_ANCHOR.TOP, line_spacing: float = 1.18):
    """文字内嵌 AUTO_SHAPE 载体（noFill 矩形，几何与 TEXT_BOX 完全一致）。

    golden-B-r211 皮肤口径：页上文字一律活在形状里（旧标杆 148 个 AUTO_SHAPE / 0 个
    TEXT_BOX），纯 TEXT_BOX 漂浮层观感呆板且不可整体拖动；可编辑性不变（形状文本框
    与文本框同一编辑模型）。
    """
    sp = _add_box(slide, x, y, w, h)
    sp.text_frame.vertical_anchor = anchor
    _write_runs(sp.text_frame, text, size=size, color=color, bold=bold, align=align,
                single_line=single_line, line_spacing=line_spacing)
    return sp


def _figure_for_assets(assets_bytes: dict[str, bytes], asset_ids: list[str]) -> tuple[str, bytes] | None:
    for aid in asset_ids:
        if aid in assets_bytes and assets_bytes[aid]:
            return aid, assets_bytes[aid]
    return None


def _plan_layout(req: PageRequest) -> tuple[list[ShapeBox], dict[str, Any]]:
    """语义 layout_hint → 引擎内部几何（内容 clamp 防溢出）。"""
    s = req.slide
    kind = s.get("layout_hint") or s.get("type") or "concept"
    title: str = s.get("title", "")
    content: str = s.get("content", "")
    asset_ids: list[str] = s.get("assets", []) or []
    has_fig = _figure_for_assets(req.assets_bytes, asset_ids) is not None
    boxes: list[ShapeBox] = []
    el: dict[str, Any] = {}

    # 通用：accent bar + 类型 chip + 页脚（QA 豁免）
    boxes.append(ShapeBox("accent", "accent", 0.60, 0.50, 0.16, 0.52, qa_exempt=True))
    boxes.append(ShapeBox("footer", "footer", 0.60, 7.06, 12.13, 0.30, req.deck_title,
                          font_pt=9, single_line=True, qa_exempt=True))
    boxes.append(ShapeBox("pageno", "footer", 12.05, 7.06, 0.68, 0.30, f"{req.page_no:02d} / {req.total_pages:02d}",
                          font_pt=9, single_line=True, qa_exempt=True))
    # chip 文本：内部枚举 → 中文标签（防英文 chip 泄漏）；缺失则用"章节"
    chip_zh = TYPE_ZH_LABEL.get(str(s.get("type", "")), "章节")
    boxes.append(ShapeBox("chip", "chip", 11.15, 0.55, 1.58, 0.42, chip_zh, font_pt=10,
                          single_line=True, qa_exempt=True))

    title_lines = M.safe_line_count(title, 10.2, 28) if len(title) > 26 else 1
    t_h = 0.52 * title_lines + 0.06
    boxes.append(ShapeBox("title", "title", 0.92, 0.44, 10.2, t_h, title, font_pt=28,
                          single_line=(title_lines == 1)))

    if kind in ("title",):
        boxes.append(ShapeBox("body", "body", 2.67, 3.30, 8.0, 1.5, content or "", font_pt=18))
        el["body"] = (2.67, 3.30, 8.0, 1.5)
        el["centered"] = True
    elif has_fig:
        fig_w, fig_h = (5.6, 4.6)
        text_w = 6.1
        body = M.clamp_lines_to_fit(content or "", text_w, 4.7, 16)
        boxes.append(ShapeBox("body", "body", 0.92, 1.55, text_w, 4.7, body, font_pt=16))
        boxes.append(ShapeBox("figure", "figure", 7.45, 1.55, fig_w, fig_h, "", font_pt=0))
        el["body"] = (0.92, 1.55, text_w, 4.7)
        el["figure"] = (7.45 + (fig_w - 5.4) / 2, 1.55 + (fig_h - 4.2) / 2, 5.4, 4.2)
    elif kind in ("definition", "concept"):
        lead = content.split("\n")[0] if content else ""
        rest = "\n".join(content.split("\n")[1:]) if "\n" in content else ""
        def_text = M.clamp_lines_to_fit(lead, 11.2, 1.75, 18)
        boxes.append(ShapeBox("defbox", "body", 0.92, 1.62, 11.2, 1.75, def_text, font_pt=18))
        if rest:
            rest_c = M.clamp_lines_to_fit(rest, 11.2, 2.9, 16)
            boxes.append(ShapeBox("body", "body", 0.92, 3.62, 11.2, 2.9, rest_c, font_pt=16))
            el["body"] = (0.92, 3.62, 11.2, 2.9)
        el["defbox"] = (0.92, 1.62, 11.2, 1.75)
    elif kind in ("derivation", "practice", "worked-example"):
        body = M.clamp_lines_to_fit(content or "", 11.4, 4.9, 16)
        boxes.append(ShapeBox("body", "body", 0.92, 1.55, 11.4, 4.9, body, font_pt=16))
        el["body"] = (0.92, 1.55, 11.4, 4.9)
        el["math_lines"] = True
    else:  # introduction / comparison / image-focus / inquiry / summary / overview
        body = M.clamp_lines_to_fit(content or "", 11.4, 4.9, 16)
        boxes.append(ShapeBox("body", "body", 0.92, 1.55, 11.4, 4.9, body, font_pt=16))
        el["body"] = (0.92, 1.55, 11.4, 4.9)
    return boxes, el


def _content_sha256(req: PageRequest, asset_hashes: dict[str, str]) -> str:
    vp = req.visual_plan or {}
    payload = {
        "builder_version": BUILDER_VERSION,
        "deck_id": req.deck_id,
        "deck_title": req.deck_title,
        "page_no": req.page_no,
        "total_pages": req.total_pages,
        "slide": req.slide,
        "asset_hashes": asset_hashes,
        # VisualPlan 进内容寻址：planner/family 变化 → 重渲染（checkpoint 语义正确）
        "plan_version": vp.get("intent", {}).get("plan_version") or (vp.get("version") if isinstance(vp.get("version"), str) else None),
        "plan_family": vp.get("layout_family"),
    }
    canon = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(canon.encode("utf-8")).hexdigest()


def build_page(req: PageRequest) -> BuiltPage:
    # R21.1：带 per-slide visual plan → family-based 渲染路径（坐标仅由 layout resolver 产生）；
    # 无 plan → legacy basic path（feature-disable 安全回退，任务书 §30）。
    if isinstance(req.visual_plan, dict) and req.visual_plan.get("layout_family"):
        return _build_page_from_plan(req)
    return _build_page_basic(req)


def _build_page_basic(req: PageRequest) -> BuiltPage:
    """legacy basic path（R21 1.0.1 形态）：无 visual_plan 时的安全回退。"""
    prs = Presentation()
    prs.slide_width = M.emu(SLIDE_W_IN)
    prs.slide_height = M.emu(SLIDE_H_IN)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    _set_bg(slide, SURFACE)

    boxes, el = _plan_layout(req)

    # 装饰：accent bar + chip 底
    _add_box(slide, 0.60, 0.50, 0.16, 0.52, fill=BRAND)
    _add_box(slide, 11.15, 0.55, 1.58, 0.42, fill=BRAND_BG, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    _add_text(slide, 11.15, 0.55, 1.58, 0.42, TYPE_ZH_LABEL.get(str(req.slide.get("type", "")), "章节"),
              size=10, color=BRAND, align=PP_ALIGN.CENTER, single_line=True, anchor=MSO_ANCHOR.MIDDLE)

    # 标题
    title = req.slide.get("title", "")
    tb = _add_text(slide, 0.92, 0.44, 10.2, boxes[[i for i, b in enumerate(boxes) if b.shape_id == "title"][0]].h,
                   title, size=28, bold=True, single_line=(len(title) <= 26))
    _ = tb

    body_geo = el.get("body")
    math_lines = el.get("math_lines", False)
    centered = el.get("centered", False)

    # 正文（分段：数学行高亮）；definition/concept 首行已进定义卡 → 正文去掉首行
    if body_geo:
        x, y, w, h = body_geo
        raw = req.slide.get("content", "") or ""
        if "defbox" in el:
            raw = "\n".join(raw.split("\n")[1:])
        lines = _norm_bullets(raw)
        if not lines and req.slide.get("content") and "defbox" not in el:
            # R21.1 缺陷修复 A：定义卡已占用首行（lead）时，fallback 严禁把含 lead 的
            # 完整原文塞回正文（尾随空行等边缘 case 曾导致 lead/body 重复）。
            lines = [req.slide["content"]]
        # 标题页正文居中
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.04)
        tf.margin_right = Inches(0.04)
        used = 0
        for ln in lines:
            if used >= 200:
                break
            is_math = math_lines and _is_math_line(ln)
            p = tf.paragraphs[0] if used == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER if centered else PP_ALIGN.LEFT
            p.line_spacing = 1.18
            if is_math:
                p.space_before = Pt(6)
                p.space_after = Pt(6)
            r = p.add_run()
            r.text = ln
            f = r.font
            f.name = FONT_PRIMARY
            f.size = Pt(18 if is_math else (18 if centered else 16))
            f.bold = is_math
            f.color.rgb = BRAND if is_math else (INK if centered else INK)
            rPr = r._r.get_or_add_rPr()
            from lxml import etree

            ea = rPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
            if ea is None:
                ea = etree.SubElement(rPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
            ea.set("typeface", FONT_PRIMARY)
            used += 1

    # 定义框（definition/concept 首段高亮卡）
    if "defbox" in el:
        x, y, w, h = el["defbox"]
        _add_box(slide, x - 0.06, y - 0.06, w + 0.12, h + 0.12, fill=BRAND_BG,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
        # 文本框画在卡片之上（z 序：后添加在上）
        lead = (req.slide.get("content", "") or "").split("\n")[0]
        _add_text(slide, x, y, w, h, lead, size=18, color=INK, anchor=MSO_ANCHOR.MIDDLE)

    # 图形/图片
    has_figure = False
    fig = _figure_for_assets(req.assets_bytes, req.slide.get("assets", []) or [])
    if fig and "figure" in el:
        aid, blob = fig
        x, y, w, h = el["figure"]
        slide.shapes.add_picture(io.BytesIO(blob), Inches(x), Inches(y), Inches(w), Inches(h))
        has_figure = True

    # 页脚
    _add_text(slide, 0.60, 7.06, 12.13, 0.30, req.deck_title, size=9, color=SUBTLE, single_line=True)
    _add_text(slide, 12.05, 7.06, 0.68, 0.30, f"{req.page_no:02d} / {req.total_pages:02d}", size=9,
              color=SUBTLE, align=PP_ALIGN.RIGHT, single_line=True)

    # 备注（含学情引用溯源，可编辑）
    notes_text = req.slide.get("speaker_note") or ""
    src_refs = req.slide.get("source_refs", []) or []
    if src_refs:
        notes_text = (notes_text + "\n" if notes_text else "") + "来源: " + " ; ".join(src_refs)
    if notes_text:
        slide.notes_slide.notes_text_frame.text = notes_text

    buf = io.BytesIO()
    prs.save(buf)
    fragment = buf.getvalue()

    asset_hashes = {aid: hashlib.sha256(b).hexdigest() for aid, b in req.assets_bytes.items()}
    all_text = title + "\n" + (req.slide.get("content", "") or "") + "\n" + notes_text
    math_chars = sorted({ch for ch in all_text if ch in MATH_CHARS})

    return BuiltPage(
        fragment_bytes=fragment,
        boxes=boxes,
        content_sha256=_content_sha256(req, asset_hashes),
        notes_text=notes_text,
        math_chars=math_chars,
        has_figure=has_figure,
    )


# ─────────── R21.1 Visual Planning Layer 渲染路径 ───────────

# ── R22 皮肤移植（golden-B-r211 口径，纯视觉层）──────────────────────────────
# 旧标杆：正文块 = 浅底圆角卡片（EFF6FF/F1F5F9）+ 文字 noFill AUTO_SHAPE 叠放；
# 行级 bullet（h 小）保持透明承载。卡片只作装饰底：qa_exempt，不进 QA 度量盒，
# 坐标仍由 layout_resolver 唯一产生（本处只做确定性视觉偏移）。
_SKIN_PAD_X = 0.14        # 卡片左右 padding（golden 实测 0.22~0.30，取近值防邻卡贴边）
_SKIN_PAD_Y_TALL = 0.10   # 大块正文卡上下 padding（行级卡不加，防行距 0.2 时贴连）
_SKIN_TALL_H = 1.10       # 高于该值视为"大块正文"（加纵向 padding）
_SKIN_CARD_MIN_H = 0.75   # 低于该高度的行级 bullet 不卡片化（golden：透明承载）
_SKIN_RADIUS = 0.08
_SKIN_CARD_EXEMPT_PREFIX = ("der-", "ann-")  # 推导/等式注释行有 dot/公式强调结构，不叠卡


def _skin_card_fill(shape_id: str) -> RGBColor:
    """卡片底色：numbered 偶数行 EFF6FF / 其余 F1F5F9（golden-B-r211 slide2 实测分布）。"""
    if shape_id.startswith("item-"):
        try:
            idx = int(shape_id.rsplit("-", 1)[1])
        except (ValueError, IndexError):
            idx = 0
        if idx % 2 == 0:
            return BRAND_BG  # EFF6FF
    if shape_id == "check-q":
        return BRAND_BG
    return SOFT_BG  # F1F5F9


def _skin_covers(rect, text_el) -> bool:
    """rect（带填充）是否完整罩住 text 元素（resolver 已发的卡：lead/def-banner/prompt 等）。"""
    return (
        rect.x <= text_el.x + 1e-6
        and rect.y <= text_el.y + 1e-6
        and rect.x + rect.w >= text_el.x + text_el.w - 1e-6
        and rect.y + rect.h >= text_el.y + text_el.h - 1e-6
    )


def _skin_cards(els: list) -> list[Element]:
    """为无卡背书的 body 文本块生成卡底 Element（只产出绘制指令，不进 QA 度量盒）。"""
    filled_rects = [e for e in els if e.kind == "rect" and e.fill]
    small_marks = [
        e for e in filled_rects if e.w <= 0.75  # 编号点 / 步骤点 / accent 小块
    ]
    cards: list[Element] = []
    for e in els:
        if e.kind != "text" or not (e.text or "").strip():
            continue
        if e.qa_kind != "body" or e.h < _SKIN_CARD_MIN_H:
            continue  # 行级 bullet：透明承载（golden 口径）
        if e.single_line:
            continue
        if any(e.shape_id.startswith(p) for p in _SKIN_CARD_EXEMPT_PREFIX):
            continue
        if any(_skin_covers(r, e) for r in filled_rects):
            continue  # resolver 已发卡（lead/def-banner/prompt/ex-prob 等）
        x, y, w, h = e.x, e.y, e.w, e.h
        # 左侧同排紧邻的编号点/accent 小块 → 卡片扩展覆盖（golden：卡包住编号圆点）
        for d in small_marks:
            if abs((d.y + d.h / 2) - (y + h / 2)) > 0.35:
                continue
            if d.x + d.w > x or x - (d.x + d.w) > 0.20:
                continue
            x, w = d.x - _SKIN_PAD_X, (e.x + e.w + _SKIN_PAD_X) - (d.x - _SKIN_PAD_X)
            break
        else:
            x, w = e.x - _SKIN_PAD_X, e.w + 2 * _SKIN_PAD_X
        pad_y = _SKIN_PAD_Y_TALL if h >= _SKIN_TALL_H else 0.0
        cards.append(Element(
            "rect", f"{e.shape_id}-card", e.qa_kind,
            x, y - pad_y, w, h + 2 * pad_y,
            fill=str(_skin_card_fill(e.shape_id)),  # RGBColor → "EFF6FF" 形式
            radius=_SKIN_RADIUS, qa_exempt=True,
        ))
    return cards


def _build_page_from_plan(req: PageRequest) -> BuiltPage:
    """family-based 渲染：ResolvedVisualPlan(语义) → layout_resolver(坐标) → OOXML。

    可编辑性不变：真实文本框/形状/图片/备注；QA 仍对 resolver 产生的盒子独立复核。
    """
    from presenton_adapter.engine import layout_resolver as LR

    vp = req.visual_plan or {}
    tokens = req.deck_style_tokens
    els = LR.resolve_page(req.page_no, req.total_pages, req.deck_title,
                          req.slide, vp, tokens)

    # TP-H1：图片字节缺失时，禁止空框伪装 —— 找出无字节的 picture，
    # 联动 drop 同一 layout_resolver 输出的配对"diag-frame"/"fig-frame"等容器 rect，
    # 并就地插入"图象暂缺"提示卡。
    fig_frame_ids_to_drop: set[str] = set()
    notice_elems: list = []
    for e in els:
        if e.kind == "picture":
            fig = _figure_for_assets(req.assets_bytes, e.asset_ids)
            if not fig:
                # 配对 frame 一并 drop（figure-frame / diag-frame / fig-frame 等约定前缀）
                fig_frame_ids_to_drop.update({
                    e.shape_id + "-frame",
                    "fig-frame",
                    "diag-frame",
                })
                # 就地插入缺图提示（覆盖 picture 的几何）
                notice_elems.extend([
                    Element("rect", "diag-missing-card", "figure",
                            e.x, e.y, e.w, e.h,
                            fill="EFF6FF", stroke="E2E8F0", radius=0.10, qa_exempt=True),
                    Element("text", "diag-missing-title", "kicker",
                            e.x + 0.30, e.y + 0.18, e.w - 0.60, 0.36,
                            text="图象暂缺", font_pt=15,
                            bold=True, color="2563EB", align="center", anchor="middle",
                            single_line=True),
                    Element("text", "diag-missing-body", "support",
                            e.x + 0.30, e.y + 0.62, e.w - 0.60, max(e.h - 0.72, 0.4),
                            text="本页为图象资产缺失降级样例。请补传教材原图或教师上传图后重渲染。",
                            font_pt=14, color="64748B", align="center", anchor="middle"),
                ])
    if fig_frame_ids_to_drop or notice_elems:
        # 重建元素列表（保序）：drop frame、drop 无字节 picture、追加 notice
        pruned: list = []
        for e in els:
            if e.kind == "rect" and e.shape_id in fig_frame_ids_to_drop:
                continue
            if e.kind == "picture" and not _figure_for_assets(req.assets_bytes, e.asset_ids):
                continue
            pruned.append(e)
        pruned.extend(notice_elems)
        els = pruned

    prs = Presentation()
    prs.slide_width = M.emu(SLIDE_W_IN)
    prs.slide_height = M.emu(SLIDE_H_IN)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    _set_bg(slide, SURFACE)

    # R22 皮肤移植：正文块卡片先画（最底层），文字随后进形状载体（golden-B-r211 口径）
    for card in _skin_cards(els):
        _add_rect_el(slide, card)

    has_figure = False
    for e in els:
        if e.kind == "rect":
            _add_rect_el(slide, e)
        elif e.kind == "text":
            _add_text_el(slide, e)
        elif e.kind == "picture":
            fig = _figure_for_assets(req.assets_bytes, e.asset_ids)
            if fig and e.w > 0.2 and e.h > 0.2:
                aid, blob = fig
                slide.shapes.add_picture(io.BytesIO(blob), Inches(e.x), Inches(e.y),
                                         Inches(e.w), Inches(e.h))
                has_figure = True

    # 备注（含溯源，可编辑）
    notes_text = req.slide.get("speaker_note") or ""
    src_refs = req.slide.get("source_refs", []) or []
    if src_refs:
        notes_text = (notes_text + "\n" if notes_text else "") + "来源: " + " ; ".join(src_refs)
    if notes_text:
        slide.notes_slide.notes_text_frame.text = notes_text

    buf = io.BytesIO()
    prs.save(buf)
    fragment = buf.getvalue()

    asset_hashes = {aid: hashlib.sha256(b).hexdigest() for aid, b in req.assets_bytes.items()}
    all_text = "\n".join(e.text for e in els if e.text) + "\n" + notes_text
    math_chars = sorted({ch for ch in all_text if ch in MATH_CHARS})

    return BuiltPage(
        fragment_bytes=fragment,
        boxes=LR.boxes_from_elements(els),
        content_sha256=_content_sha256(req, asset_hashes),
        notes_text=notes_text,
        math_chars=math_chars,
        has_figure=has_figure,
    )


def _hex(color: str) -> RGBColor:
    color = color.lstrip("#")
    return RGBColor(int(color[0:2], 16), int(color[2:4], 16), int(color[4:6], 16))


def _add_rect_el(slide, e) -> None:
    shape = MSO_SHAPE.ROUNDED_RECTANGLE if e.radius is not None else MSO_SHAPE.RECTANGLE
    fill = _hex(e.fill) if e.fill else None
    stroke = _hex(e.stroke) if e.stroke else None
    _add_box(slide, e.x, e.y, e.w, e.h, fill=fill, line=stroke, shape=shape, radius=e.radius)


def _add_text_el(slide, e) -> None:
    align = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[e.align]
    anchor = MSO_ANCHOR.MIDDLE if e.anchor == "middle" else MSO_ANCHOR.TOP
    # R22 皮肤：文字进 AUTO_SHAPE 载体（noFill，几何不变），页上零裸 TEXT_BOX
    tb = _add_shape_text(slide, e.x, e.y, e.w, e.h, e.text, size=e.font_pt, color=_hex(e.color),
                         bold=e.bold, align=align, single_line=e.single_line, anchor=anchor,
                         line_spacing=e.line_spacing)
    _ = tb
